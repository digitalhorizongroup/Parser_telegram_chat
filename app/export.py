import logging
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill

from app.config import EXCEL_FILE
from app.models import ChatResult

logger = logging.getLogger(__name__)

HEADERS = ["Рейтинг", "Название", "Telegram", "ID", "Участники", "Рост", "Можно писать",
           "Последняя активность", "24ч", "7 дней", "Среднее/7д", "30 дней", "Среднее/30д",
           "Активные авторы", "Просмотры", "Реакции", "ER, %", "Онлайн", "Язык", "Голосовой чат", "Обсуждение"]


def _safe(value):
    return f"'{value}" if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value


def _row(item: ChatResult):
    return [item.score, _safe(item.title), item.url, item.chat_id, item.participants, item.audience_growth,
            "Да" if item.can_write else "Нет", item.last_activity, item.messages_24h, item.messages_7d,
            item.average_7d, item.messages_30d, item.average_30d, item.active_authors, item.views,
            item.reactions, item.engagement_rate, item.online_users, item.language,
            item.voice_type or "Нет", item.discussion_url]


def _prepare_sheet(ws, items: list[ChatResult]):
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FFC000")
        cell.alignment = Alignment(horizontal="center")
    for item in sorted(items, key=lambda x: x.score, reverse=True):
        ws.append(_row(item))
        for column in (3, 21):
            cell = ws.cell(ws.max_row, column)
            if cell.value:
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:U{max(1, ws.max_row)}"
    ws.conditional_formatting.add(f"I2:I{max(2, ws.max_row)}", ColorScaleRule(
        start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50,
        mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    widths = [12, 34, 28, 14, 14, 12, 16, 25] + [13] * 13
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width


def save_to_excel(data: list[ChatResult], filename=EXCEL_FILE, errors=None, run_id=None):
    filename = Path(filename)
    filename.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = "Сводка"
    chats = [item for item in data if item.chat_type == "chat"]
    channels = [item for item in data if item.chat_type == "channel"]
    summary.append(["Отчёт Telegram Parser", datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S")])
    summary.append(["Запуск", run_id or "-"])
    summary.append(["Чаты", len(chats)])
    summary.append(["Каналы", len(channels)])
    summary.append(["Ошибки/фильтры", len(errors or [])])
    summary.append(["Сообщения за 24ч", sum(x.messages_24h for x in data)])
    chart = BarChart()
    chart.title = "Результаты"
    chart.add_data(Reference(summary, min_col=2, min_row=3, max_row=5))
    chart.set_categories(Reference(summary, min_col=1, min_row=3, max_row=5))
    summary.add_chart(chart, "D2")
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 22

    _prepare_sheet(wb.create_sheet("Чаты"), chats)
    _prepare_sheet(wb.create_sheet("Каналы"), channels)
    err = wb.create_sheet("Ошибки")
    err.append(["ID", "Название", "Тип", "Сообщение", "Время"])
    for item in errors or []:
        err.append([item.get("chat_id"), _safe(item.get("title")), item.get("kind"), _safe(item.get("message")), item.get("occurred_at")])
    err.freeze_panes = "A2"
    err.auto_filter.ref = f"A1:E{max(1, err.max_row)}"
    wb.save(filename)
    logger.info("Отчёт сохранён: %s", filename)
